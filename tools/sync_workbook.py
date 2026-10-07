"""Synchronize the frozen experiment workbook with the runtime material design."""

from __future__ import annotations

import os
from collections import Counter
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.workbook.properties import CalcProperties

from uav_sim.config import load_experiment_data


WORKBOOK = Path("data/UAV_Trial_Design.xlsx")
CANDIDATE = Path("data/.UAV_Trial_Design.syncing.xlsx")


def populated_formulas(workbook):
    return {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    }


def style_snapshot(workbook):
    return {
        (sheet.title, cell.coordinate): copy(cell._style)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def structure_snapshot(workbook):
    return {
        "sheetnames": list(workbook.sheetnames),
        "merges": {
            sheet.title: tuple(sorted(str(item) for item in sheet.merged_cells.ranges))
            for sheet in workbook.worksheets
        },
        "table_names": {
            sheet.title: tuple(sorted(sheet.tables)) for sheet in workbook.worksheets
        },
        "freeze_panes": {
            sheet.title: str(sheet.freeze_panes) if sheet.freeze_panes else None
            for sheet in workbook.worksheets
        },
    }


def copy_row_style(sheet, source_row: int, target_row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        target._style = copy(source._style)
        if source.has_style:
            target.number_format = source.number_format
            target.alignment = copy(source.alignment)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def main() -> None:
    data = load_experiment_data(WORKBOOK, redesign_wind_zones=True)
    workbook = load_workbook(WORKBOOK, data_only=False)
    before_formulas = populated_formulas(workbook)
    before_styles = style_snapshot(workbook)
    before_structure = structure_snapshot(workbook)

    readme = workbook["README"]
    readme_updates = {
        "B2": "UAV 监督决策实验冻结材料；调度目标为最小化预计飞行时间，地面等待单独记录",
        "B3": "Predictability(H/M/L) × Intervenability(H/L)；被试内 3×2，每格 3 trial",
        "B4": "所有 trial 仅显示 1 架无人机",
        "B5": "release / delay 120 s / reroute；三类最优行动各 6 trial",
        "B6": "Path-A 直达；Path-B 北绕；正式判断题明确限定参考路径与行动",
        "B7": "顺风/侧风/逆风参考夹角与最优行动交叉轮换；T0 为实测值，T1/T2 显示预测范围",
        "B8": "正式评分、练习未来仿真和过程题统一使用图中平均风场；固定 seed 的隐藏实现仅用于开发检验",
        "B9": "Intervenability=(最差-最优)/最差：High 35%–50%，Low 5%–15%；最优—次优差值单独平衡",
        "B10": "版本 4.1 / 2026-08-pretest-flow-v2；18 个正式 trial，3 block × 6 trial",
    }
    for coordinate, value in readme_updates.items():
        readme[coordinate] = value

    trials = workbook["Trials"]
    for row_index, trial_id in enumerate(sorted(data.trials), start=2):
        trial = data.trials[trial_id]
        values = (
            trial.trial_id,
            trial.block,
            trial.predictability,
            trial.intervenability,
            trial.complexity,
            trial.drone_count,
            trial.route_set_id,
            trial.start_node,
            trial.goal_node,
            trial.release_path,
            trial.delay_seconds,
            trial.reroute_path,
            trial.planned_best_action,
            trial.intermediate_question,
            trial.seed,
        )
        for column, value in enumerate(values, start=1):
            trials.cell(row_index, column).value = value
    final_trial_row = len(data.trials) + 1
    if trials.max_row > final_trial_row:
        trials.delete_rows(final_trial_row + 1, trials.max_row - final_trial_row)

    routes = workbook["Routes"]
    path_id_column = next(
        cell.column for cell in routes[1] if cell.value == "pathId"
    )
    for row_index in range(routes.max_row, 1, -1):
        if routes.cell(row_index, path_id_column).value not in {"Path-A", "Path-B"}:
            routes.delete_rows(row_index, 1)
    routes.tables["RoutesTable"].ref = f"A1:F{routes.max_row}"
    trials.tables["TrialsTable"].ref = f"A1:O{final_trial_row}"

    truth_sheet = workbook["ActionGroundTruth"]
    truth_headers = (
        "trialId", "predictability", "intervenability", "sceneLoad", "routeSetId",
        "releaseTime_s", "delayTime_s", "rerouteTime_s", "releaseEnergy_Wh",
        "delayEnergy_Wh", "rerouteEnergy_Wh", "bestTimeAction", "bestEnergyAction",
        "decisionGap_pct", "intervenability_pct", "designCheck",
    )
    copy_row_style(truth_sheet, 1, 1, len(truth_headers))
    for column, header in enumerate(truth_headers, start=1):
        truth_sheet.cell(1, column).value = header
    truth_sheet.cell(1, 16)._style = copy(truth_sheet.cell(1, 15)._style)
    for row_index, trial_id in enumerate(sorted(data.trials), start=2):
        trial = data.trials[trial_id]
        truth = data.ground_truth[trial_id]
        inputs = (
            trial_id,
            trial.predictability,
            trial.intervenability,
            trial.complexity,
            trial.route_set_id,
            truth.release_time_s,
            truth.delay_time_s,
            truth.reroute_time_s,
            truth.release_energy_wh,
            truth.delay_energy_wh,
            truth.reroute_energy_wh,
        )
        for column, value in enumerate(inputs, start=1):
            truth_sheet.cell(row_index, column).value = value
        copy_row_style(truth_sheet, row_index, row_index, len(truth_headers))
        truth_sheet.cell(row_index, 12).value = (
            f'=IF(F{row_index}<=MIN(G{row_index},H{row_index}),"release",'
            f'IF(G{row_index}<=H{row_index},"delay","reroute"))'
        )
        truth_sheet.cell(row_index, 13).value = (
            f'=IF(I{row_index}<=MIN(J{row_index},K{row_index}),"release",'
            f'IF(J{row_index}<=K{row_index},"delay","reroute"))'
        )
        truth_sheet.cell(row_index, 14).value = (
            f'=(MEDIAN(F{row_index}:H{row_index})-MIN(F{row_index}:H{row_index}))/'
            f'MIN(F{row_index}:H{row_index})'
        )
        truth_sheet.cell(row_index, 15).value = (
            f'=(MAX(F{row_index}:H{row_index})-MIN(F{row_index}:H{row_index}))/'
            f'MAX(F{row_index}:H{row_index})'
        )
        truth_sheet.cell(row_index, 16).value = (
            f'=IF(AND(N{row_index}>=5%,N{row_index}<=15%),'
            f'IF(C{row_index}="H",IF(AND(O{row_index}>=35%,O{row_index}<=50%),"PASS","RECALIBRATE"),'
            f'IF(AND(O{row_index}>=5%,O{row_index}<=15%),"PASS","RECALIBRATE")),"RECALIBRATE")'
        )
        truth_sheet.cell(row_index, 14).number_format = "0.0%"
        truth_sheet.cell(row_index, 15).number_format = "0.0%"
        truth_sheet.cell(row_index, 16)._style = copy(truth_sheet.cell(row_index, 15)._style)
    final_truth_row = len(data.trials) + 1
    if truth_sheet.max_row > final_truth_row:
        truth_sheet.delete_rows(final_truth_row + 1, truth_sheet.max_row - final_truth_row)
    truth_sheet.tables["GroundTruthTable"].ref = f"A1:P{final_truth_row}"

    wind_sheet = workbook["WindZones"]
    wind_rows = []
    for trial_id in sorted(data.trials):
        for layer in (0, 1, 2):
            for zone in data.wind_zones[trial_id][layer]:
                wind_rows.append((
                    trial_id,
                    layer,
                    zone.zone_id,
                    zone.cx,
                    zone.cy,
                    zone.radius,
                    zone.mean_speed,
                    zone.direction_deg,
                    zone.speed_sd,
                    zone.direction_sd,
                ))
    for row_index, values in enumerate(wind_rows, start=2):
        for column, value in enumerate(values, start=1):
            wind_sheet.cell(row_index, column).value = value
    final_wind_row = len(wind_rows) + 1
    if wind_sheet.max_row > final_wind_row:
        wind_sheet.delete_rows(final_wind_row + 1, wind_sheet.max_row - final_wind_row)
    wind_sheet.tables["WindZonesTable"].ref = f"A1:J{final_wind_row}"

    parameters = workbook["Parameters"]
    legacy_parameters = {
        "highIntervenabilityMinGap": "deprecated_bestSecondMinGap_v3",
        "lowIntervenabilityMaxGap": "deprecated_bestSecondMaxGap_v3",
    }
    for row in range(2, parameters.max_row + 1):
        name = parameters.cell(row, 1).value
        if name in legacy_parameters:
            parameters.cell(row, 1).value = legacy_parameters[name]
            parameters.cell(row, 2).value = "unused in v4"
    parameter_updates = {
        "vAir_mps": 12,
        "minGroundSpeed_mps": 2,
        "maxGroundSpeed_mps": 18,
        "windUpdate_s": 0.5,
        "Pbase_W": 95,
        "headwindCoeff_W_per_mps2": 2.2,
        "crosswindCoeff_W_per_mps2": 0.8,
        "payloadCoeff_W_per_kg": 12,
        "batteryCapacity_Wh": 90,
        "reserveFraction": 0.2,
        "delaySeconds": 120,
        "highPredictabilitySpeedSD_mps": 0.5,
        "mediumPredictabilitySpeedSD_mps": 1.0,
        "lowPredictabilitySpeedSD_mps": 1.5,
        "highPredictabilityDirectionSD_deg": 12,
        "mediumPredictabilityDirectionSD_deg": 35,
        "lowPredictabilityDirectionSD_deg": 60,
        "windDirectionConvention": "vector-to: 0° east, 90° north",
        "windLayer1Start_s": 120,
        "windLayer2Start_s": 240,
        "highIntervenabilityMinPct": 0.35,
        "highIntervenabilityMaxPct": 0.50,
        "lowIntervenabilityMinPct": 0.05,
        "lowIntervenabilityMaxPct": 0.15,
        "decisionGapMinPct": 0.05,
        "decisionGapMaxPct": 0.15,
        "plannedEffectiveParticipants": 100,
        "experimentVersion": "4.1",
        "materialVersion": "2026-08-pretest-flow-v2",
    }
    existing_rows = {
        parameters.cell(row, 1).value: row for row in range(2, parameters.max_row + 1)
    }
    for parameter, value in parameter_updates.items():
        row = existing_rows.get(parameter)
        if row is None:
            row = parameters.max_row + 1
            copy_row_style(parameters, 17, row, 2)
            parameters.cell(row, 1).value = parameter
            existing_rows[parameter] = row
        parameters.cell(row, 2).value = value
        if parameter.endswith("Pct"):
            parameters.cell(row, 2).number_format = "0%"
    parameters.tables["ParametersTable"].ref = f"A1:B{parameters.max_row}"

    workbook.calculation = CalcProperties(
        calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True,
    )
    workbook.save(CANDIDATE)

    checked = load_workbook(CANDIDATE, data_only=False)
    after_formulas = populated_formulas(checked)
    after_structure = structure_snapshot(checked)
    if len(after_formulas) != 90:
        raise AssertionError(f"Expected 90 formulas, found {len(after_formulas)}")
    for formula in after_formulas.values():
        Tokenizer(formula)
        if any(error in formula for error in ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!")):
            raise AssertionError(f"Formula error token found: {formula}")
    if before_structure["sheetnames"] != after_structure["sheetnames"]:
        raise AssertionError("Worksheet order changed")
    for key in ("merges", "table_names", "freeze_panes"):
        if before_structure[key] != after_structure[key]:
            raise AssertionError(f"Workbook structure changed unexpectedly: {key}")
    for key, style in before_styles.items():
        sheet_name, coordinate = key
        row_number, _ = coordinate_to_tuple(coordinate)
        if sheet_name == "WindZones" and row_number > 109:
            continue
        if sheet_name == "Trials" and row_number > 19:
            continue
        if sheet_name == "ActionGroundTruth" and row_number > 19:
            continue
        if sheet_name == "ActionGroundTruth" and coordinate[0] in {"N", "O", "P"}:
            continue
        if sheet_name == "Routes":
            continue
        if checked[sheet_name][coordinate]._style != style:
            raise AssertionError(f"Existing style changed: {sheet_name}!{coordinate}")

    if checked["WindZones"].max_row != 109:
        raise AssertionError("WindZones must contain 108 data rows")
    if checked["WindZones"].tables["WindZonesTable"].ref != "A1:J109":
        raise AssertionError("WindZones table range was not updated")
    if checked["Trials"].tables["TrialsTable"].ref != "A1:O19":
        raise AssertionError("Trials table range was not updated")
    if checked["ActionGroundTruth"].tables["GroundTruthTable"].ref != "A1:P19":
        raise AssertionError("ActionGroundTruth table range was not updated")
    if checked["Routes"].tables["RoutesTable"].ref != "A1:F43":
        raise AssertionError("Routes table must contain only Path-A and Path-B")
    checked_parameters = {
        checked["Parameters"].cell(row, 1).value: checked["Parameters"].cell(row, 2).value
        for row in range(2, checked["Parameters"].max_row + 1)
    }
    expected_predictability_parameters = {
        "highPredictabilitySpeedSD_mps": 0.5,
        "mediumPredictabilitySpeedSD_mps": 1.0,
        "lowPredictabilitySpeedSD_mps": 1.5,
        "highPredictabilityDirectionSD_deg": 12,
        "mediumPredictabilityDirectionSD_deg": 35,
        "lowPredictabilityDirectionSD_deg": 60,
    }
    for parameter, expected in expected_predictability_parameters.items():
        if checked_parameters.get(parameter) != expected:
            raise AssertionError(f"{parameter} was not synchronized")

    trial_rows = list(checked["Trials"].iter_rows(min_row=2, values_only=True))
    cell_counts = Counter((row[2], row[3]) for row in trial_rows)
    if len(cell_counts) != 6 or set(cell_counts.values()) != {3}:
        raise AssertionError(f"Invalid 3x2 trial cells: {cell_counts}")
    if {row[4] for row in trial_rows} != {"Fixed"}:
        raise AssertionError("Scene load must be fixed, not an experimental factor")
    if {row[5] for row in trial_rows} != {1}:
        raise AssertionError("All trials must display only the UAV")
    for block in (1, 2, 3):
        block_cells = {(row[2], row[3]) for row in trial_rows if row[1] == block}
        if len(block_cells) != 6:
            raise AssertionError(f"Block {block} does not contain all 6 cells")
    if Counter(row[12] for row in trial_rows) != Counter({
        "release": 6, "delay": 6, "reroute": 6,
    }):
        raise AssertionError("Planned best actions are not balanced")

    os.replace(CANDIDATE, WORKBOOK)
    print({
        "workbook": str(WORKBOOK.resolve()),
        "sheets": checked.sheetnames,
        "formula_count": len(after_formulas),
        "trial_count": len(trial_rows),
        "factorial_cells": {"-".join(key): value for key, value in sorted(cell_counts.items())},
        "wind_zone_rows": checked["WindZones"].max_row - 1,
        "parameter_rows": checked["Parameters"].max_row - 1,
        "table_refs": {
            sheet.title: {name: sheet.tables[name].ref for name in sheet.tables}
            for sheet in checked.worksheets
            if sheet.tables
        },
    })


if __name__ == "__main__":
    main()
